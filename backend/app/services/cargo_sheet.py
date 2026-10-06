"""내 화물 현황 표와 엑셀 받기 (019 마이그레이션, 10-06 정운 결정, 아버지가 손으로 쓰는 관리표 기준).

책임: 화물 표(user_cargo)와 신용장 표(user_lc)를 읽어 "신용장 한 묶음에 B/L 여러 줄" 표로 묶고, 신용장·출항 칸을 쓰고, .xlsx로 만든다.
  load_sheet()      묶음 목록. 신용장 있는 묶음은 개설일 순, 신용장 없는 줄은 마지막 한 묶음 (GET /api/cargo/sheet)
  set_lc()          화물 한 건을 신용장에 잇거나 푼다. 신용장은 번호로 찾아 없으면 만든다 (PUT .../lc)
  set_departed()    출항일을 쓴다. 사용자가 넣으면 manual, 도착 통지서 사진에서 읽으면 document (PUT .../departed)
  build_xlsx()      내려받는 엑셀 파일 바이트 (GET /api/cargo/sheet.xlsx)

칸의 출처 (019): 입항(arrived_on)·입고(warehoused_on)는 조회 때 UNI-PASS 이력에서 코드가 채운다(services/cargo.stage_date).
신용장과 출항은 UNI-PASS·해수부에 없어서 사용자가 넣는다. 이 모듈이 쓰는 칸은 그 둘뿐이다.
LLM은 이 표를 읽지도 쓰지도 않는다. DB 연결은 부르는 쪽이 연다. 여기 함수는 conn만 받아 SQL을 보낸다.
"""
from __future__ import annotations

import io
import unicodedata
from datetime import date, datetime, timezone

from backend.app import crypto
from backend.app.db import atomic
from backend.app.services.evidence import KST

# 신용장 번호 글자 수 상한. 은행이 매기는 번호는 SWIFT 참조 번호(최대 16자) 안팎이라 넉넉히 35로 둔다 (10-06).
# 같은 숫자가 019 마이그레이션 CHECK에도 있다
LC_NO_MAX = 35
# 한 번에 읽는 화물 줄 상한. 화면에 한꺼번에 그릴 수 있는 크기로 정했다 (화물 표 all_rows는 100)
SHEET_MAX = 500
SOURCES = ("manual", "document")

# 엑셀의 B/L 번호 칸에 무엇을 적을지 (10-06). 정운이 아직 정하지 않아 기본은 가린 값("masked")이다.
# "original"로 바꾸면 잠가 둔 원본(bl_encrypted)을 열어 적는다. 바꾸는 곳은 이 상수 하나뿐이다.
# 원본을 파일로 내려보내면 파일이 사용자 PC·메일에 남으므로 SDD 3.5(가린 값만 남긴다)와 어긋나는 선택이다. 정해지면 SDD를 먼저 고친다
XLSX_BL_SOURCE = "masked"

# 엑셀 머리글. 앞 여섯 칸은 아버지 표 그대로이고 뒤 셋은 UNI-PASS에서 오는 참고 칸이다 (10-06)
HEADERS = ("신용장 개설일", "신용장 번호", "B/L 번호", "출항", "입항", "입고", "포워더", "선사", "컨테이너 번호")
WIDTHS = (14, 22, 20, 10, 12, 10, 24, 20, 26)
DATE_FORMAT = "m/d"  # 아버지 표처럼 월/일만. 엑셀에는 진짜 날짜로 들어가서 정렬·계산이 된다
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class CargoNotFound(Exception):
    """그 사용자의 화물 줄이 없다. 다른 사람의 화물 지문도 같은 결과다(있는지 알려 주지 않는다)."""


def clean_lc_no(value: str) -> str:
    """신용장 번호를 다듬어 돌려준다. 앞뒤 공백을 지우고, 비었거나 35자를 넘거나 제어 문자(줄바꿈·탭 등)가 있으면 ValueError.
    제어 문자를 막는 이유: 이 값은 엑셀 칸과 화면에 그대로 나가므로 줄바꿈·이상한 글자로 표를 깨뜨리지 못하게 한다."""
    text = value.strip()
    if not text:
        raise ValueError("신용장 번호가 비어 있습니다")
    if len(text) > LC_NO_MAX:
        raise ValueError(f"신용장 번호는 {LC_NO_MAX}자 이하여야 합니다")
    if any(unicodedata.category(ch) in ("Cc", "Cf", "Zl", "Zp") for ch in text):
        raise ValueError("신용장 번호에 쓸 수 없는 글자가 있습니다")
    return text


# 읽어 올 칸. 맨 뒤 셋(lc 쪽)은 LEFT JOIN이라 신용장이 없으면 None이다
_SELECT = """SELECT c.cargo_key, c.bl_masked, c.bl_encrypted, c.label, c.status, c.departed_on, c.departed_source,
                    c.arrived_on, c.last_eta, c.warehoused_on, c.last_stage, c.forwarder_name, c.carrier_name,
                    c.container_no, c.created_at, l.id, l.lc_no, l.opened_on, l.created_at
             FROM user_cargo c LEFT JOIN user_lc l ON l.id = c.lc_id
             WHERE c.user_id = %s ORDER BY c.created_at LIMIT %s"""
_FIELDS = ("cargo_key", "bl_masked", "bl_encrypted", "label", "status", "departed_on", "departed_source",
           "arrived_on", "eta", "warehoused_on", "last_stage", "forwarder_name", "carrier_name",
           "container_no", "created_at")
# API로 내보내는 칸. bl_encrypted·created_at 같은 안쪽 값은 넣지 않는다 (API 계약 10-06)
_ROW_OUT = tuple(f for f in _FIELDS if f not in ("bl_encrypted", "created_at"))


def _grouped(conn, user_id: str) -> list[dict]:
    """줄을 신용장별로 묶는다. 안쪽 값(bl_encrypted, created_at)을 아직 지니고 있다. 엑셀이 원본 B/L을 열 때 쓴다.
    - 신용장 묶음: 개설일 빠른 순, 개설일 없는 것은 뒤, 같으면 신용장을 만든 순
    - 묶음 안 줄: 화물을 만든 순 (SQL이 이미 created_at 순이고 파이썬 정렬은 안정이라 순서가 유지된다)
    - 신용장 없는 줄은 맨 끝 묶음 하나 (lc가 None)"""
    by_lc: dict[int, dict] = {}
    loose: list[dict] = []
    for r in conn.execute(_SELECT, [user_id, SHEET_MAX]).fetchall():
        row = dict(zip(_FIELDS, r[:len(_FIELDS)]))
        lc_id, lc_no, opened_on, lc_created = r[len(_FIELDS):]
        if lc_id is None:
            loose.append(row)
            continue
        group = by_lc.setdefault(lc_id, {"lc": {"id": lc_id, "lc_no": lc_no, "opened_on": opened_on},
                                          "rows": [], "_created": lc_created})
        group["rows"].append(row)
    ordered = sorted(by_lc.values(), key=lambda g: (g["lc"]["opened_on"] is None, g["lc"]["opened_on"] or date.min,
                                                    g["_created"] or datetime.min.replace(tzinfo=timezone.utc)))
    groups = [{"lc": g["lc"], "rows": g["rows"]} for g in ordered]
    if loose:
        groups.append({"lc": None, "rows": loose})
    return groups


def load_sheet(conn, user_id: str) -> dict:
    """API 응답 모양의 표 {"groups": [...]}. 안쪽 값은 걸러 낸다. 신용장 묶음이 없고 줄도 없으면 groups는 빈 목록이다."""
    return {"groups": [{"lc": g["lc"], "rows": [{k: row[k] for k in _ROW_OUT} for row in g["rows"]]}
                       for g in _grouped(conn, user_id)]}


def _owned(conn, user_id: str, key: str) -> None:
    """그 사용자의 화물 줄인지 확인한다. 아니면 CargoNotFound. 다른 사용자의 지문과 없는 지문을 구별해 주지 않는다."""
    if conn.execute("SELECT 1 FROM user_cargo WHERE user_id = %s AND cargo_key = %s", [user_id, key]).fetchone() is None:
        raise CargoNotFound(key)


def set_lc(conn, user_id: str, key: str, lc_no: str | None, opened_on: date | None) -> None:
    """화물 한 건을 신용장에 잇는다. lc_no가 None이면 신용장에서 푼다 (PUT .../lc).

    - 신용장은 (사용자, 번호)로 하나다. 있으면 개설일만 새로 들어왔을 때 덮고, 안 들어왔으면 기존 값을 둔다
      (같은 번호를 두 번째 B/L에 이으면서 개설일을 또 적지 않아도 되게)
    - 풀거나 옮긴 뒤 줄이 하나도 없는 신용장은 지운다. 빈 신용장은 표에 안 보이니 남겨 둘 이유가 없다
    - 신용장 만들기·잇기·정리는 한 트랜잭션이다. 중간에 실패해도 반쪽 상태가 남지 않는다
    lc_no는 부르는 쪽이 clean_lc_no로 다듬어 넘긴다."""
    _owned(conn, user_id, key)
    with atomic(conn):
        lc_id = None
        if lc_no is not None:
            lc_id = conn.execute(
                """INSERT INTO user_lc (user_id, lc_no, opened_on) VALUES (%s, %s, %s)
                   ON CONFLICT (user_id, lc_no) DO UPDATE SET
                     opened_on = COALESCE(EXCLUDED.opened_on, user_lc.opened_on), updated_at = now()
                   RETURNING id""", [user_id, lc_no, opened_on]).fetchone()[0]
        conn.execute("UPDATE user_cargo SET lc_id = %s, updated_at = now() WHERE user_id = %s AND cargo_key = %s",
                     [lc_id, user_id, key])
        conn.execute("""DELETE FROM user_lc WHERE user_id = %s
                        AND NOT EXISTS (SELECT 1 FROM user_cargo WHERE user_cargo.lc_id = user_lc.id)""", [user_id])


def set_departed(conn, user_id: str, key: str, departed_on: date | None, source: str = "manual") -> None:
    """출항일을 쓴다 (PUT .../departed). 날짜를 비우면(None) 출처도 같이 비운다.
    source: manual은 사용자가 직접 넣은 값, document는 포워더 도착 통지서 사진에서 읽어 사용자가 확인한 값 (019).
    출항은 UNI-PASS·해수부에 없어서 이 두 길뿐이다."""
    if source not in SOURCES:
        raise ValueError(f"source는 {SOURCES} 중 하나여야 합니다")
    _owned(conn, user_id, key)
    conn.execute("UPDATE user_cargo SET departed_on = %s, departed_source = %s, updated_at = now() "
                 "WHERE user_id = %s AND cargo_key = %s",
                 [departed_on, source if departed_on else None, user_id, key])


# ── 엑셀 ──

def bl_for_xlsx(row: dict) -> str:
    """엑셀 B/L 번호 칸에 적을 글자. XLSX_BL_SOURCE가 "original"이면 원본, 아니면 가린 값이다.
    원본을 열지 못하면(열쇠가 없거나 잘못됨) 가린 값으로 대신한다. 빈 칸으로 내보내지는 않는다."""
    if XLSX_BL_SOURCE == "original" and row.get("bl_encrypted"):
        return crypto.decrypt(row["bl_encrypted"]) or row["bl_masked"]
    return row["bl_masked"]


def _text(cell, value: str) -> None:
    """글자 칸을 쓴다. 값이 =로 시작해도 수식으로 풀리지 않도록 글자 형식(s)으로 못 박는다.
    신용장 번호는 사용자 입력이고 포워더·선사·컨테이너는 외부(UNI-PASS) 값이라, 수식 주입으로 파일을 여는 사람의 엑셀이 움직이면 안 된다."""
    cell.value = value
    cell.data_type = "s"


def build_xlsx(conn, user_id: str) -> bytes:
    """내 화물 현황 표를 .xlsx 바이트로 만든다 (GET /api/cargo/sheet.xlsx).

    칸 순서는 HEADERS. 신용장 개설일·번호는 아버지 표처럼 묶음의 첫 줄에만 적는다.
    날짜는 진짜 날짜 칸(m/d 모양)이다. 입항일이 없고 예정일(eta)만 있으면 글자로 "예정 m/d"를 적는다. 실제 신고된 날과 섞이지 않게 한다.
    머리글은 굵게, 첫 줄은 고정, 칸 너비는 WIDTHS. 줄이 하나도 없어도 머리글만 있는 파일을 돌려준다."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "내 화물 현황"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for letter, width in zip("ABCDEFGHI", WIDTHS):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"

    for group in _grouped(conn, user_id):
        lc = group["lc"]
        for i, row in enumerate(group["rows"]):
            r = ws.max_row + 1
            if lc and i == 0:
                if lc["opened_on"]:
                    ws.cell(r, 1, lc["opened_on"]).number_format = DATE_FORMAT
                _text(ws.cell(r, 2), lc["lc_no"])
            _text(ws.cell(r, 3), bl_for_xlsx(row))
            for col, field in ((4, "departed_on"), (5, "arrived_on"), (6, "warehoused_on")):
                if row[field]:
                    ws.cell(r, col, row[field]).number_format = DATE_FORMAT
            if not row["arrived_on"] and row["eta"]:
                _text(ws.cell(r, 5), f"예정 {row['eta'].month}/{row['eta'].day}")
            for col, field in ((7, "forwarder_name"), (8, "carrier_name"), (9, "container_no")):
                if row[field]:
                    _text(ws.cell(r, col), row[field])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def xlsx_filename(now: datetime | None = None) -> str:
    """내려받는 파일 이름 marvia_cargo_YYYYMMDD.xlsx. 날짜는 한국 날짜다."""
    return f"marvia_cargo_{(now or datetime.now(timezone.utc)).astimezone(KST):%Y%m%d}.xlsx"
